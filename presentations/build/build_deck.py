#!/usr/bin/env python3
"""
Build presentations/presentation.pptx

Content source : distilled-project/presentation/project-presentation-content.md
Format source  : distilled-project/presentation/EL-phase-1.md  (EL Phase-I deck)

Design language reproduced from the EL Phase-I reference:
  - numbered section narrative (Introduction -> ... -> References)
  - a dedicated outline slide listing every section
  - bold lead-in labels on bullets ("Topic: ... text")
  - a dedicated flowchart slide
  - a phase-based timeline slide
  - an IEEE-style references slide
  - an "expected outcome by end of first year" slide
  - an image-only closing slide

Run:  /tmp/pptxenv/bin/python presentations/build/build_deck.py
"""

import os
import math
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
ASSETS = os.path.join(ROOT, "presentations", "assets")
OUT = os.path.join(ROOT, "presentations", "presentation.pptx")
os.makedirs(ASSETS, exist_ok=True)

# ----------------------------------------------------------------------------
# Palette / typography
# ----------------------------------------------------------------------------
NAVY = RGBColor(0x12, 0x26, 0x3A)
INK = RGBColor(0x16, 0x20, 0x2B)
BODY = RGBColor(0x44, 0x4E, 0x59)
MUTED = RGBColor(0x87, 0x91, 0x9D)
RULE = RGBColor(0xD9, 0xDF, 0xE6)
PANEL = RGBColor(0xF3, 0xF6, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACC = RGBColor(0x1F, 0x6F, 0xB2)      # steel blue  - primary accent
ACC2 = RGBColor(0xE8, 0x73, 0x4A)     # amber       - emphasis / predicted
GOOD = RGBColor(0x2E, 0x8B, 0x6F)     # green       - supported
WARN = RGBColor(0xB8, 0x8A, 0x1E)     # amber-dark  - caveats
ACC_SOFT = RGBColor(0xE7, 0xF0, 0xF8)
ACC2_SOFT = RGBColor(0xFD, 0xEE, 0xE7)
PANEL2 = RGBColor(0xFA, 0xFB, 0xFD)

FONT = "Calibri"
MONO = "Consolas"

W, H = 13.333, 7.5
M = 0.72                 # side margin
CW = W - 2 * M           # content width  = 11.893
TOP = 1.52               # content top
BOT = 6.86               # content bottom

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]

PAGE = {"n": 0}


# ----------------------------------------------------------------------------
# Primitives
# ----------------------------------------------------------------------------
def E(inches):
    return Inches(inches)


def est(lines, width, size, line=1.19, sa=5.0):
    """Conservative rendered-height estimate (inches) for wrapped text."""
    if isinstance(lines, str):
        lines = [lines]
    cpl = max(1, int(width / (size * 0.46 / 72.0)))
    total = 0.0
    for ln in lines:
        n = max(1, math.ceil(len(ln) / cpl)) if ln.strip() else 1
        total += n * (size * 1.19 / 72.0) * line + sa / 72.0
    return total * 1.07


def rect(s, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE, adj=None):
    sh = s.shapes.add_shape(shape, E(x), E(y), E(w), E(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    if adj is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = adj
        except Exception:
            pass
    sh.text_frame.text = ""
    return sh


def tb(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = s.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, first=False, align=PP_ALIGN.LEFT, space_after=6, line=1.15, space_before=0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    p.line_spacing = line
    return p


def run(p, text, size=14, bold=False, color=INK, font=FONT, italic=False):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return r


def label(s, x, y, w, text, size=11, color=MUTED, bold=True, font=FONT):
    """Small uppercase section kicker."""
    tf = tb(s, x, y, w, 0.28)
    run(para(tf, first=True), text.upper(), size=size, bold=bold, color=color, font=font)
    return tf


def bullet(s, x, y, w, lead, text, size=14, space_after=11, marker="▪", mcolor=ACC):
    """EL convention: bold lead-in label, then the point."""
    tf = tb(s, x, y, w, 0.6)
    p = para(tf, first=True, space_after=space_after)
    run(p, marker + "  ", size=size, bold=True, color=mcolor)
    if lead:
        run(p, lead + " ", size=size, bold=True, color=INK)
    run(p, text, size=size, bold=False, color=BODY)
    return tf


def plain(s, x, y, w, text, size=14, color=BODY, bold=False, align=PP_ALIGN.LEFT, line=1.25):
    tf = tb(s, x, y, w, 0.6)
    run(para(tf, first=True, align=align, line=line), text, size=size, color=color, bold=bold)
    return tf


def card(s, x, y, w, h, title=None, body=None, title_size=14.5, body_size=12.5,
         accent=ACC, fill=PANEL, title_color=None, line=None, adj=0.05):
    """Card with a coloured spine. Grows automatically so text never spills out."""
    body = [] if body is None else (body if isinstance(body, list) else [body])
    inner_w = w - 0.45
    need = 0.16 + (0.44 if title else 0) + (est(body, inner_w, body_size, sa=5) if body else 0) + 0.16
    h = max(h, need)
    sh = rect(s, x, y, w, h, fill=fill, line=line or RULE, lw=0.9,
              shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=adj)
    rect(s, x, y, 0.055, h, fill=accent)
    ty = y + 0.16
    if title:
        tf = tb(s, x + 0.26, ty, inner_w, 0.4)
        run(para(tf, first=True, space_after=4), title, size=title_size, bold=True,
            color=title_color or INK)
        ty += 0.44
    if body:
        tf = tb(s, x + 0.26, ty, inner_w, h - (ty - y) - 0.14)
        for i, ln in enumerate(body):
            run(para(tf, first=(i == 0), space_after=5, line=1.19), ln,
                size=body_size, color=BODY)
    return sh


def callout(s, x, y, w, h, title, text, accent=WARN, fill=RGBColor(0xFD, 0xF7, 0xE6),
            tsize=12.5, bsize=12.5):
    """Amber caveat box. Grows automatically so text never spills out."""
    body = text if isinstance(text, list) else [text]
    inner_w = w - 0.46
    need = (0.17 + est([title], inner_w, tsize, sa=4)
            + est(body, inner_w, bsize, sa=3, line=1.17) + 0.17)
    h = max(h, need)
    rect(s, x, y, w, h, fill=fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.04)
    rect(s, x, y, 0.06, h, fill=accent)
    tf = tb(s, x + 0.24, y + 0.15, inner_w, h - 0.30)
    run(para(tf, first=True, space_after=4), title, size=tsize, bold=True, color=INK)
    for ln in body:
        run(para(tf, space_after=3, line=1.17), ln, size=bsize, color=BODY)
    return h


def table(s, x, y, w, rows, cols, data, col_w=None, row_h=0.42, head_h=0.44,
          fsize=11.5, hsize=11.5, head_fill=NAVY, head_color=WHITE, zebra=True,
          align=None, bold_col0=True):
    gt = s.shapes.add_table(rows, cols, E(x), E(y), E(w), E(row_h * rows))
    tbl = gt.table
    tbl.first_row = True
    tbl.horz_banding = False
    if col_w:
        total = sum(col_w)
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = E(w * cw / total)
    tbl.rows[0].height = E(head_h)
    for r in range(1, rows):
        tbl.rows[r].height = E(row_h)
    for r in range(rows):
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.margin_left = E(0.11)
            cell.margin_right = E(0.11)
            cell.margin_top = E(0.05)
            cell.margin_bottom = E(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = head_fill
            else:
                cell.fill.fore_color.rgb = PANEL2 if (zebra and r % 2 == 0) else WHITE
            tf = cell.text_frame
            tf.word_wrap = True
            txt = data[r][c]
            bold = (r == 0) or (bold_col0 and c == 0 and r > 0)
            al = PP_ALIGN.LEFT
            if align and align[c] == "c":
                al = PP_ALIGN.CENTER
            run(para(tf, first=True, align=al, space_after=0, line=1.12), txt,
                size=(hsize if r == 0 else fsize), bold=bold,
                color=(head_color if r == 0 else INK))
    return gt


def code_panel(s, x, y, w, h, title, lines, accent=ACC, tsize=11.5, csize=10.5):
    rect(s, x, y, w, h, fill=RGBColor(0xF7, 0xF9, 0xFC), line=RULE, lw=0.9,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.03)
    rect(s, x, y, w, 0.34, fill=RGBColor(0xE9, 0xEE, 0xF4))
    tf = tb(s, x + 0.16, y + 0.055, w - 0.3, 0.26)
    run(para(tf, first=True), title, size=tsize, bold=True, color=INK)
    tf = tb(s, x + 0.18, y + 0.46, w - 0.34, h - 0.58)
    for i, ln in enumerate(lines):
        col = BODY
        if ln.strip().startswith("//"):
            col = MUTED
        run(para(tf, first=(i == 0), space_after=1, line=1.02), ln,
            size=csize, color=col, font=MONO)
    return None


def arrow_down(s, x, y, h, color=ACC, w=0.16):
    sh = rect(s, x - w / 2, y, w, h, fill=color, shape=MSO_SHAPE.DOWN_ARROW)
    return sh


def chip(s, x, y, w, h, text, fill=ACC, color=WHITE, size=11.5, bold=True):
    sh = rect(s, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.35)
    tf = sh.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = 0
    run(para(tf, first=True, align=PP_ALIGN.CENTER, space_after=0), text,
        size=size, bold=bold, color=color)
    return sh


def stat(s, x, y, w, h, value, caption, accent=ACC2):
    rect(s, x, y, w, h, fill=PANEL, line=RULE, lw=0.9,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.06)
    tf = tb(s, x + 0.12, y + 0.16, w - 0.24, 0.72)
    run(para(tf, first=True, align=PP_ALIGN.CENTER, space_after=0), value,
        size=32, bold=True, color=accent)
    tf = tb(s, x + 0.14, y + 0.92, w - 0.28, h - 1.0)
    for i, ln in enumerate(caption if isinstance(caption, list) else [caption]):
        run(para(tf, first=(i == 0), align=PP_ALIGN.CENTER, space_after=2, line=1.15),
            ln, size=11, color=BODY)


# ----------------------------------------------------------------------------
# Slide frames
# ----------------------------------------------------------------------------
def content_slide(title, kicker, section=None, subtitle=None):
    PAGE["n"] += 1
    s = prs.slides.add_slide(BLANK)
    label(s, M, 0.40, CW, kicker, size=11, color=ACC)
    tf = tb(s, M, 0.66, CW, 0.58)
    run(para(tf, first=True, space_after=0), title, size=26, bold=True, color=INK)
    rect(s, M, 1.30, 1.15, 0.055, fill=ACC)
    if subtitle:
        tf = tb(s, M, 1.42, CW, 0.3)
        run(para(tf, first=True), subtitle, size=12.5, color=MUTED)
    # footer
    tf = tb(s, M, 7.03, 8.0, 0.28)
    run(para(tf, first=True), (section or kicker).title(), size=9.5, color=MUTED)
    tf = tb(s, W - M - 1.2, 7.03, 1.2, 0.28)
    run(para(tf, first=True, align=PP_ALIGN.RIGHT), str(PAGE["n"]), size=9.5, color=MUTED)
    return s


def divider(num, title, subtitle):
    PAGE["n"] += 1
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, W, H, fill=NAVY)
    rect(s, 0, 0, 0.22, H, fill=ACC)
    tf = tb(s, 1.05, 2.28, 3.2, 1.5)
    run(para(tf, first=True, space_after=0), num, size=74, bold=True, color=ACC2)
    tf = tb(s, 1.05, 3.62, 10.6, 0.85)
    run(para(tf, first=True, space_after=0), title, size=38, bold=True, color=WHITE)
    tf = tb(s, 1.05, 4.62, 10.4, 0.9)
    run(para(tf, first=True, line=1.25), subtitle, size=15, color=RGBColor(0xB9, 0xC6, 0xD4))
    rect(s, 1.05, 4.42, 1.5, 0.055, fill=ACC2)
    return s


# ----------------------------------------------------------------------------
# Figure: the three hypotheses (schematic, no measured values)
# ----------------------------------------------------------------------------
def make_hypotheses_figure(path):
    fig, ax = plt.subplots(figsize=(7.5, 3.9), dpi=220)
    xs = [0, 1, 2]
    labels = ["Random", "BM25", "Semantic"]
    h1 = [0.30, 0.55, 0.74]
    h0a = [0.56, 0.575, 0.565]
    h0b = [0.30, 0.66, 0.44]

    ax.plot(xs, h0a, "-", color="#9AA5B1", lw=2.2, marker="o", ms=7, label="H0a  flat — strategy makes no difference")
    ax.plot(xs, h1, "--", color="#1F6FB2", lw=2.2, marker="s", ms=6.5, label="H1  monotone — harder is always better")
    ax.plot(xs, h0b, "-", color="#E8734A", lw=3.4, marker="o", ms=9, label="H0b  inverted-U — helps, then hurts  (predicted)")

    ax.set_xticks(xs)
    ax.set_xticklabels(labels, fontsize=13, color="#16202B")
    ax.set_yticks([])
    ax.set_xlim(-0.32, 2.32)
    ax.set_ylim(0.14, 0.88)
    ax.set_ylabel("Embedding quality (schematic)", fontsize=11.5, color="#444E59")
    ax.set_xlabel("Hardness of the negative", fontsize=11.5, color="#444E59")

    for sp in ("top", "right", "left"):
        ax.spines[sp].set_visible(False)
    ax.spines["bottom"].set_color("#D9DFE6")
    ax.tick_params(axis="x", length=0, pad=8)
    ax.grid(axis="y", visible=False)

    ax.annotate("peak", xy=(1, 0.66), xytext=(1.30, 0.79), fontsize=11,
                color="#C0602F", fontweight="bold",
                arrowprops=dict(arrowstyle="->", color="#C0602F", lw=1.4))

    leg = ax.legend(loc="lower left", fontsize=10.5, frameon=False, bbox_to_anchor=(-0.03, -0.06))
    for t in leg.get_texts():
        t.set_color("#444E59")

    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


# ============================================================================
# SLIDES
# ============================================================================

# ---------------------------------------------------------------- 1  TITLE
PAGE["n"] += 1
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, W, H, fill=NAVY)
rect(s, 0, 0, 0.28, H, fill=ACC)
rect(s, 7.55, 0, 5.78, H, fill=RGBColor(0x0E, 0x1E, 0x2E))

label(s, 1.0, 1.28, 6.4, "EL Phase-I  ·  Project Presentation", size=11.5, color=ACC2)
tf = tb(s, 1.0, 1.72, 6.9, 2.0)
run(para(tf, first=True, space_after=8, line=1.05), "Domain-Specific", size=31, bold=True, color=WHITE)
run(para(tf, space_after=8, line=1.05), "Contrastive Embeddings", size=31, bold=True, color=WHITE)
run(para(tf, space_after=0, line=1.05), "for Code Similarity Detection", size=31, bold=True, color=ACC2)
rect(s, 1.0, 3.98, 1.6, 0.06, fill=ACC2)
tf = tb(s, 1.0, 4.20, 6.2, 0.9)
run(para(tf, first=True, space_after=6, line=1.2),
    "Comparing negative-sampling strategies in contrastive fine-tuning",
    size=16, color=RGBColor(0xC9, 0xD6, 0xE4))
run(para(tf, space_after=0, line=1.2),
    "Random  ·  BM25  ·  Semantic", size=14, bold=True, color=RGBColor(0x9F, 0xB4, 0xC8))

rect(s, 1.0, 5.42, 6.2, 0.035, fill=RGBColor(0x2A, 0x3D, 0x52))
tf = tb(s, 1.0, 5.62, 6.2, 1.0)
run(para(tf, first=True, space_after=4, line=1.2),
    "Theme:  Contrastive representation learning  ·  Domain adaptation  ·  Software engineering",
    size=12, color=RGBColor(0x9F, 0xB4, 0xC8))
run(para(tf, space_after=0, line=1.2),
    "Test domain: Java code clone detection (BigCloneBench)   ·   Base model: GraphCodeBERT",
    size=12, color=RGBColor(0x9F, 0xB4, 0xC8))

# right-hand schematic: the triplet idea
cx = 10.44
rect(s, 8.15, 1.55, 4.35, 4.4, fill=RGBColor(0x14, 0x29, 0x40), line=RGBColor(0x24, 0x3C, 0x55), lw=1)
tf = tb(s, 8.42, 1.82, 3.8, 0.3)
run(para(tf, first=True), "THE ONE VARIABLE", size=10.5, bold=True, color=ACC2)
tf = tb(s, 8.42, 2.14, 3.8, 0.3)
run(para(tf, first=True), "Which non-clone do we push away?", size=12.5, color=RGBColor(0xB9, 0xC6, 0xD4))

rect(s, 8.42, 2.72, 3.8, 0.78, fill=RGBColor(0x1B, 0x33, 0x4D), shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.12)
tf = tb(s, 8.58, 2.86, 3.5, 0.5)
run(para(tf, first=True, space_after=2), "Anchor", size=12, bold=True, color=WHITE)
run(para(tf, space_after=0), "bubbleSort(int[] a)", size=10.5, color=RGBColor(0x9F, 0xB4, 0xC8), font=MONO)

rect(s, 8.42, 3.68, 1.82, 0.7, fill=RGBColor(0x1B, 0x4D, 0x3C), shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.14)
tf = tb(s, 8.56, 3.80, 1.55, 0.5)
run(para(tf, first=True, space_after=2), "+  Positive", size=11.5, bold=True, color=RGBColor(0x7E, 0xD9, 0xB8))
run(para(tf, space_after=0), "pull closer", size=10.5, color=RGBColor(0x9F, 0xB4, 0xC8))

rect(s, 10.40, 3.68, 1.82, 0.7, fill=RGBColor(0x4D, 0x26, 0x1B), shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.14)
tf = tb(s, 10.54, 3.80, 1.55, 0.5)
run(para(tf, first=True, space_after=2), "–  Negative", size=11.5, bold=True, color=RGBColor(0xF0, 0xA9, 0x86))
run(para(tf, space_after=0), "push away", size=10.5, color=RGBColor(0x9F, 0xB4, 0xC8))

rect(s, 9.62, 4.44, 0.42, 0.26, fill=RGBColor(0x2A, 0x3D, 0x52), shape=MSO_SHAPE.DOWN_ARROW)
rect(s, 9.62, 4.44, 0.42, 0.26, fill=RGBColor(0x2A, 0x3D, 0x52))

for i, (lab, col) in enumerate([("Random", RGBColor(0x9F, 0xB4, 0xC8)),
                                ("BM25", RGBColor(0x9F, 0xB4, 0xC8)),
                                ("Semantic", ACC2)]):
    x = 8.42 + i * 1.29
    rect(s, x, 4.80, 1.18, 0.62, fill=RGBColor(0x1B, 0x33, 0x4D), shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.16)
    tf = tb(s, x, 4.94, 1.18, 0.4)
    run(para(tf, first=True, align=PP_ALIGN.CENTER, space_after=0), lab, size=10.5, bold=True, color=col)

tf = tb(s, 1.0, 6.62, 6.2, 0.4)
run(para(tf, first=True), "Four conditions · one variable · generalisation tested from the start",
    size=11.5, color=RGBColor(0x7E, 0x92, 0xA8))

# ---------------------------------------------------------------- 2  OUTLINE
s = content_slide("Presentation Outline", "Overview", "Outline")
items = [
    ("01", "Introduction", "Why domain-specific similarity is hard"),
    ("02", "Literature Review", "Five studies, the trends, the gap"),
    ("03", "Problem Definition", "What exactly is unclear"),
    ("04", "Objectives", "Main aim and four specific objectives"),
    ("05", "Methodology", "Design, data, analysis, tools"),
    ("06", "Proposed Outcomes", "Expected results and impact"),
    ("07", "Timeline of the Project", "Initiation → closure"),
    ("08", "Conclusion", "Summary and next steps"),
    ("09", "References", "IEEE-style sources"),
]
cw, ch, gapx, gapy = 3.76, 1.52, 0.30, 0.28
for i, (num, name, sub) in enumerate(items):
    r, c = divmod(i, 3)
    x = M + c * (cw + gapx)
    y = TOP + 0.15 + r * (ch + gapy)
    rect(s, x, y, cw, ch, fill=PANEL, line=RULE, lw=0.9, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.07)
    rect(s, x, y, 0.055, ch, fill=ACC)
    tf = tb(s, x + 0.24, y + 0.17, 0.8, 0.5)
    run(para(tf, first=True), num, size=24, bold=True, color=ACC2)
    tf = tb(s, x + 0.24, y + 0.66, cw - 0.45, 0.4)
    run(para(tf, first=True, space_after=3), name, size=14.5, bold=True, color=INK)
    tf = tb(s, x + 0.24, y + 1.00, cw - 0.45, 0.45)
    run(para(tf, first=True, line=1.15), sub, size=11, color=BODY)
plain(s, M, 6.62, CW,
      "An EL Phase-I deliverable slide and a closing slide follow the references.",
      size=11, color=MUTED)

# ---------------------------------------------------------------- DIVIDER 1
divider("01", "Introduction",
        "What the topic is, and why it is worth a project.")

# ---------------------------------------------------------------- 3  OVERVIEW
s = content_slide("Overview of the Topic", "01  ·  Introduction", "Introduction",
                  "Generic embeddings are a compromise; contrastive learning removes the compromise — "
                  "but introduces a choice nobody has isolated.")
bullet(s, M, TOP + 0.02, CW,
       "The problem with generic models:",
       "Pretrained embedding models are trained on broad, general-purpose objectives. They do not "
       "inherently capture what “similar” means inside one specialised domain.", size=14.5)
bullet(s, M, TOP + 0.78, CW,
       "Contrastive learning:",
       "Train embeddings by showing pairs of similar and dissimilar examples — pulling the similar "
       "ones together and pushing the dissimilar ones apart in vector space.", size=14.5)
bullet(s, M, TOP + 1.54, CW,
       "The underexplored lever:",
       "“Dissimilar” is our choice. Does picking negatives at random, by keyword overlap, or by "
       "model similarity change the result? That is this project.", size=14.5)
bullet(s, M, TOP + 2.30, CW,
       "Our test case:",
       "Code similarity detection — a domain with strong public benchmarks, so the comparison is "
       "rigorous and the numbers are comparable.", size=14.5)

# pull / push schematic
by = TOP + 3.06
rect(s, M, by, CW, 1.42, fill=RGBColor(0xF7, 0xFA, 0xFD), line=RULE, lw=0.9,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.05)
tf = tb(s, M + 0.3, by + 0.16, 3.0, 0.3)
run(para(tf, first=True), "THE MECHANISM IN ONE PICTURE", size=10.5, bold=True, color=ACC)
# anchor dot
rect(s, M + 1.55, by + 0.75, 0.30, 0.30, fill=ACC, shape=MSO_SHAPE.OVAL)
tf = tb(s, M + 1.30, by + 1.12, 0.85, 0.3)
run(para(tf, first=True, align=PP_ALIGN.CENTER), "anchor", size=10.5, bold=True, color=INK)
# positive
rect(s, M + 3.35, by + 0.62, 0.26, 0.26, fill=GOOD, shape=MSO_SHAPE.OVAL)
tf = tb(s, M + 3.02, by + 0.98, 0.92, 0.3)
run(para(tf, first=True, align=PP_ALIGN.CENTER), "positive", size=10.5, color=GOOD, bold=True)
rect(s, M + 1.98, by + 0.86, 1.32, 0.07, fill=GOOD)
tf = tb(s, M + 2.20, by + 0.60, 1.0, 0.25)
run(para(tf, first=True, align=PP_ALIGN.CENTER), "pull together", size=10, color=GOOD)
# negative
rect(s, M + 6.35, by + 0.62, 0.26, 0.26, fill=ACC2, shape=MSO_SHAPE.OVAL)
tf = tb(s, M + 6.02, by + 0.98, 0.92, 0.3)
run(para(tf, first=True, align=PP_ALIGN.CENTER), "negative", size=10.5, color=ACC2, bold=True)
rect(s, M + 1.98, by + 1.36, 4.28, 0.07, fill=ACC2)
tf = tb(s, M + 3.55, by + 1.44, 1.4, 0.25)
run(para(tf, first=True, align=PP_ALIGN.CENTER), "push apart", size=10, color=ACC2)
tf = tb(s, M + 7.55, by + 0.62, 4.0, 0.9)
run(para(tf, first=True, space_after=4, line=1.15),
    "Everything in this project is about the red arrow.", size=13, bold=True, color=INK)
run(para(tf, space_after=0, line=1.15),
    "There are thousands of possible negatives. We compare three ways of choosing one.",
    size=11.5, color=BODY)

# ---------------------------------------------------------------- 4  IMPORTANCE
s = content_slide("Importance of the Topic", "01  ·  Introduction", "Introduction")
stat(s, M, TOP + 0.05, 3.72, 1.92, "8M+",
     ["validated code-clone pairs in BigCloneBench,",
      "drawn from ~25,000 real", "open-source Java projects"], accent=ACC)
stat(s, M + 3.94, TOP + 0.05, 3.72, 1.92, "43",
     ["distinct functionalities in the benchmark —",
      "and most clone pairs belong to",
      "only a handful of them"], accent=ACC2)
stat(s, M + 7.88, TOP + 0.05, 3.72, 1.92, "3",
     ["ways to choose a negative, all three used",
      "in practice, and almost never", "compared against each other"], accent=GOOD)

bullet(s, M, TOP + 2.20, CW,
       "It is not only about code:",
       "Domain-specific similarity is needed for legal clause matching, duplicate bug reports, "
       "medical record matching — anywhere a generic, off-the-shelf embedding is a compromise.", size=13.5)
bullet(s, M, TOP + 2.74, CW,
       "The lever is known, but under-studied here:",
       "Negative-sampling strategy is a recognised lever in contrastive learning generally. Its "
       "specific effect for domain adaptation is rarely compared systematically.", size=13.5)
bullet(s, M, TOP + 3.28, CW,
       "Why now:",
       "AI-assisted generation makes superficially different but functionally equivalent content "
       "cheap to produce — so capturing true semantic similarity matters more, not less.", size=13.5)
callout(s, M, TOP + 3.92, CW, 0.78, "Justification",
        "Better embeddings for specialised domains are not a marginal convenience: similarity "
        "search is the retrieval layer that downstream tools are built on.")

# ---------------------------------------------------------------- DIVIDER 2
divider("02", "Literature Review",
        "Five studies, the trends they establish, and the gap they leave open.")

# ---------------------------------------------------------------- 5  KEY STUDIES
s = content_slide("Key Studies and Findings", "02  ·  Literature Review", "Literature Review",
                  "Five significant studies that frame the project.")
rows = [
    ["Study", "Contribution"],
    ["Guo et al. — GraphCodeBERT (ICLR 2021)",
     "Introduced structural, data-flow-aware pretraining for code, improving on purely token-based models."],
    ["Jain et al. — Contrastive Code Representation Learning (EMNLP 2021)",
     "Applied contrastive learning directly to source code representation, demonstrating its viability for this modality."],
    ["Sonnekalb et al. — Generalizability of Code Clone Detection on CodeBERT (ASE 2022)",
     "Found that a ~96.5% F1 score on standard benchmarks drops significantly when evaluated on unseen functionalities."],
    ["Liu et al. — ContraBERT (ICSE 2023)",
     "Showed contrastive learning improves robustness of code pretrained models beyond standard fine-tuning."],
    ["Khajezade, Fard & Shehata (Empirical Software Engineering 2024)",
     "Directly evaluated contrastive learning under limited-data and unseen-problem conditions for clone detection."],
]
table(s, M, TOP + 0.12, CW, 6, 2, rows, col_w=[2.05, 3.0], row_h=0.66, head_h=0.42,
      fsize=12, hsize=12)
callout(s, M, TOP + 3.98, CW, 0.85, "The thread that connects them",
        "Each study establishes that contrastive learning helps code models, or that benchmark scores "
        "overstate generalisation. None of them isolates how the negative examples were chosen.",
        accent=ACC)

# ---------------------------------------------------------------- 6  TRENDS
s = content_slide("Trends in Literature", "02  ·  Literature Review", "Literature Review")
bullet(s, M, TOP + 0.10, CW, "Trend 1 —",
       "A shift from generic-purpose embeddings toward domain-adapted representations trained with "
       "contrastive objectives (Jain et al., 2021; Liu et al., 2023).", size=14)
bullet(s, M, TOP + 0.72, CW, "Trend 2 —",
       "Increasing scrutiny of generalisability claims: several 2022–2024 studies question whether "
       "strong benchmark scores reflect genuine understanding or benchmark-specific exposure.", size=14)
bullet(s, M, TOP + 1.34, CW, "Trend 3 —",
       "Contrastive learning is increasingly favoured over classification-based fine-tuning, because "
       "it produces reusable, general-purpose embeddings rather than a task-specific head.", size=14)
bullet(s, M, TOP + 1.96, CW, "Trend 4 —",
       "Evaluation is moving beyond a single held-out split toward held-out categories, "
       "functionalities and domains.", size=14)
rect(s, M, TOP + 2.72, CW, 1.34, fill=ACC2_SOFT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.05)
rect(s, M, TOP + 2.72, 0.06, 1.34, fill=ACC2)
tf = tb(s, M + 0.28, TOP + 2.90, CW - 0.55, 1.0)
run(para(tf, first=True, space_after=5), "The gap observed across this literature", size=14.5,
    bold=True, color=INK)
run(para(tf, space_after=0, line=1.22),
    "Negative-sampling strategy is acknowledged as influential in contrastive learning generally, but is "
    "rarely isolated and compared as an explicit experimental variable — especially for code.",
    size=13, color=BODY)
callout(s, M, TOP + 4.28, CW, 0.9, "Honest framing",
        "The three strategies are compared thoroughly in natural-language retrieval (AugSBERT; ANCE; "
        "STAR/ADORE; SimANS). What is missing is a controlled comparison for code, and — critically — "
        "any test of whether the strategy changes generalisation.")

# ---------------------------------------------------------------- 7  THEORETICAL FRAMEWORK
s = content_slide("Theoretical Framework", "02  ·  Literature Review", "Literature Review")
card(s, M, TOP + 0.08, 5.80, 2.42,
     title="Contrastive Representation Learning",
     body=["Le-Khac, Healy & Smeaton, IEEE Access, 2020.",
           "",
           "Embeddings are learned by pulling similar pairs together and pushing dissimilar pairs "
           "apart in vector space — rather than fitting fixed category labels.",
           "",
           "Consequence: the pairs are the curriculum."],
     accent=ACC, body_size=12.5)
card(s, M + 6.09, TOP + 0.08, 5.80, 2.42,
     title="Classical IR Theory — BM25",
     body=["A keyword-relevance ranking model.",
           "",
           "Used here as one method for constructing “hard” negative training pairs: retrieve the "
           "snippets that share the most keywords with the anchor, then use them as negatives.",
           "",
           "Cheap, deterministic, and blind to meaning."],
     accent=ACC2, body_size=12.5)
callout(s, M, TOP + 2.72, CW, 1.02, "Relevance to this project",
        "Similarity is learned relative to explicitly chosen pairs. That makes the choice of pairs — "
        "and specifically how the negatives are chosen — a first-class design decision, not an "
        "implementation detail. This is the theoretical basis for treating negative-sampling strategy "
        "as an independent variable.",
        accent=ACC)
tf = tb(s, M, TOP + 3.98, CW, 0.9)
run(para(tf, first=True, space_after=6, line=1.2),
    "In one sentence:", size=13.5, bold=True, color=INK)
run(para(tf, space_after=0, line=1.2),
    "If the curriculum is the pairs, then the negatives are half the curriculum — and this project is "
    "the first, in this domain, to vary only that half and measure what changes.",
    size=13, color=BODY)

# ---------------------------------------------------------------- 8  RELEVANCE
s = content_slide("Relevance to Current Research", "02  ·  Literature Review", "Literature Review")
bullet(s, M, TOP + 0.10, CW, "Comparability:",
       "We adopt the same base model (GraphCodeBERT) and the same benchmark (BigCloneBench) used by "
       "Guo et al. (2021), Sonnekalb et al. (2022) and Khajezade et al. (2024), so our numbers sit "
       "directly alongside published ones.", size=13.5)
bullet(s, M, TOP + 0.78, CW, "Extending Sonnekalb et al. (2022):",
       "They raised generalisation as a critique. We make generalisation testing a built-in, core "
       "experiment rather than a follow-up observation.", size=13.5)
bullet(s, M, TOP + 1.46, CW, "Extending Khajezade et al. (2024) and Liu et al. (2023):",
       "They show contrastive learning helps. We isolate negative-sampling strategy as the single "
       "variable under study.", size=13.5)
bullet(s, M, TOP + 2.14, CW, "Extending the retrieval literature:",
       "Random vs BM25 vs semantic negatives is a well-populated comparison in text retrieval. It has "
       "not been run for code.", size=13.5)

rect(s, M, TOP + 2.94, CW, 1.42, fill=PANEL, line=RULE, lw=0.9,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.05)
tf = tb(s, M + 0.3, TOP + 3.10, CW - 0.6, 1.16)
run(para(tf, first=True, space_after=6, line=1.2),
    "The positioning, stated carefully", size=14, bold=True, color=INK)
run(para(tf, space_after=0, line=1.22),
    "This is not a claim that nothing similar has been done. Contrastive learning for code, and "
    "generalisation gaps for code models, are both established. The specific comparison — three "
    "negative-sampling strategies, held constant everywhere else, evaluated on both a standard split "
    "and held-out functionalities — is what this project adds.",
    size=12.5, color=BODY)

# ---------------------------------------------------------------- 9  ALREADY EXISTS
s = content_slide("What Already Exists vs. Our Contribution", "02  ·  Literature Review",
                  "Literature Review")
rows = [
    ["Already exists", "Our contribution"],
    ["Contrastive learning improves code representations generally (Jain et al., 2021; Liu et al., 2023)",
     "We isolate negative-sampling strategy as the variable — random vs BM25 vs semantic, compared directly"],
    ["Generalisation gaps are documented for fine-tuned code models (Sonnekalb et al., 2022; Khajezade et al., 2024)",
     "We build generalisation testing into the core experimental design, not a separate critique"],
    ["Domain-specific fine-tuning is studied per domain, in isolation",
     "We frame it as a method-level question — does negative strategy matter? — using code as one well-supported test case"],
]
table(s, M, TOP + 0.12, CW, 4, 2, rows, col_w=[1.0, 1.0], row_h=1.05, head_h=0.44,
      fsize=12.5, hsize=12.5)
callout(s, M, TOP + 3.88, CW, 0.92, "Scope discipline",
        "We do not claim the comparison is new in general — it is well established for natural-language "
        "retrieval. We claim it is unrun for code, and that the interaction between negative strategy and "
        "generalisation has not been measured in any domain.",
        accent=ACC)

# ---------------------------------------------------------------- DIVIDER 3
divider("03", "Problem Definition",
        "What, precisely, is unclear — and what it costs to leave it unclear.")

# ---------------------------------------------------------------- 10  PROBLEM STATEMENT
s = content_slide("Clear Statement of the Problem", "03  ·  Problem Definition", "Problem Definition")
rect(s, M, TOP + 0.06, CW, 1.62, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.045)
rect(s, M, TOP + 0.06, 0.07, 1.62, fill=ACC2)
tf = tb(s, M + 0.42, TOP + 0.30, CW - 0.85, 1.2)
run(para(tf, first=True, line=1.28, space_after=0),
    "“It is unclear whether the method used to select negative training examples during contrastive "
    "fine-tuning meaningfully affects the quality and generalization of domain-specific embeddings, "
    "or whether such improvements — if any — hold up beyond the training distribution.”",
    size=15.5, bold=False, color=WHITE)
bullet(s, M, TOP + 1.92, CW, "Context 1 —",
       "Generic pretrained models are trained on broad objectives such as masked language modelling, "
       "not explicitly for domain-specific similarity judgement.", size=13.5)
bullet(s, M, TOP + 2.50, CW, "Context 2 —",
       "Contrastive fine-tuning is a known adaptation method, but the choice of negative pairs is "
       "rarely treated as a controlled variable in published studies.", size=13.5)
bullet(s, M, TOP + 3.08, CW, "Context 3 —",
       "Prior work shows benchmark performance can substantially overstate real generalisation, which "
       "raises the question of whether that depends on how the model was trained.", size=13.5)
callout(s, M, TOP + 3.78, CW, 1.0, "The two questions hidden inside the problem statement",
        ["1.  Does the negative-selection method change embedding quality at all?",
         "2.  Does it change how well that quality survives on data unlike the training data?"],
        accent=ACC2)

# ---------------------------------------------------------------- 11  SIGNIFICANCE
s = content_slide("Significance of Addressing the Problem", "03  ·  Problem Definition",
                  "Problem Definition")
card(s, M, TOP + 0.06, 3.80, 2.28, "If unaddressed",
     ["Practitioners adapting embeddings to a new domain default to the simplest option — random "
      "negatives — without evidence that it is the best choice."],
     accent=ACC2, body_size=12.5)
card(s, M + 4.04, TOP + 0.06, 3.80, 2.28, "The gap",
     ["Without a systematic comparison, real and available performance improvements from better "
      "negative sampling may be left unclaimed."],
     accent=ACC, body_size=12.5)
card(s, M + 8.09, TOP + 0.06, 3.80, 2.28, "The consequence",
     ["Deployed similarity systems underperform silently on real-world, unseen inputs while appearing "
      "strong on standard benchmarks."],
     accent=WARN, body_size=12.5)

rect(s, M, TOP + 2.56, CW, 1.30, fill=ACC_SOFT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.05)
rect(s, M, TOP + 2.56, 0.06, 1.30, fill=ACC2)
tf = tb(s, M + 0.3, TOP + 2.74, CW - 0.6, 1.0)
run(para(tf, first=True, space_after=5, line=1.2),
    "The size of the prize, measured by others", size=14, bold=True, color=INK)
run(para(tf, space_after=0, line=1.22),
    "Re-evaluated on functionality categories the model has never seen, task-specific clone detectors "
    "lose up to 48% F1, averaging 31% — against benchmark scores above 95%. A benchmark number that "
    "high is not measuring what a practitioner needs it to measure.",
    size=12.5, color=BODY)

callout(s, M, TOP + 4.06, CW, 1.16, "Why this is the right moment to ask",
        "Negative-sampling strategy is cheap to change — it is a data-construction choice, not a new "
        "architecture or a bigger model. If it matters, it is one of the cheapest gains available. If it "
        "does not, that is worth knowing before the field spends more effort on elaborate mining schemes.",
        accent=GOOD)

# ---------------------------------------------------------------- DIVIDER 4
divider("04", "Objectives",
        "One main objective, four specific ones, and what the project contributes.")

# ---------------------------------------------------------------- 12  OBJECTIVES
s = content_slide("Objectives", "04  ·  Objectives", "Objectives")
rect(s, M, TOP + 0.02, CW, 1.06, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.06)
rect(s, M, TOP + 0.02, 0.07, 1.06, fill=ACC2)
tf = tb(s, M + 0.4, TOP + 0.17, CW - 0.8, 0.8)
run(para(tf, first=True, space_after=4), "Main objective", size=11, bold=True, color=ACC2)
run(para(tf, space_after=0, line=1.2),
    "To evaluate whether the choice of negative-sampling strategy in contrastive fine-tuning "
    "meaningfully affects domain-specific embedding quality and generalization, using code similarity "
    "detection as a controlled test case.",
    size=13.5, color=WHITE)

objs = [
    ("1", "Fine-tune GraphCodeBERT using three distinct negative-sampling strategies: random, "
          "BM25-based and semantically-mined.", ACC),
    ("2", "Compare the embedding quality of each strategy against an off-the-shelf, non-fine-tuned "
          "baseline.", ACC),
    ("3", "Evaluate whether performance improvements generalise to code functionalities not "
          "represented during training.", ACC2),
    ("4", "Demonstrate differences in embedding quality across strategies, visually and empirically.",
     GOOD),
]
for i, (n, txt, col) in enumerate(objs):
    y = TOP + 1.26 + i * 0.66
    rect(s, M, y, CW, 0.58, fill=PANEL, line=RULE, lw=0.9,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.09)
    chip(s, M + 0.16, y + 0.10, 0.36, 0.36, n, fill=col, size=12)
    tf = tb(s, M + 0.70, y + 0.10, CW - 0.9, 0.42)
    p = para(tf, first=True, space_after=0, line=1.15)
    run(p, txt, size=13, color=BODY)

# ---------------------------------------------------------------- 13  CONTRIBUTIONS
s = content_slide("Expected Contributions", "04  ·  Objectives", "Objectives")
card(s, M, TOP + 0.08, 3.80, 2.62, "A controlled comparison",
     ["A direct comparison of negative-sampling strategies for contrastive fine-tuning in this domain, "
      "with everything else held constant.",
      "",
      "The strategy comparison is established for text retrieval; it has not been run for code."],
     accent=ACC, body_size=12.5)
card(s, M + 4.04, TOP + 0.08, 3.80, 2.62, "Cost-aware guidance",
     ["Evidence on whether “harder” mining strategies justify their added implementation cost over "
      "simple random sampling.",
      "",
      "Practically useful: mining pipelines cost engineering time, and today that decision is made on "
      "intuition."],
     accent=GOOD, body_size=12.5)
card(s, M + 8.09, TOP + 0.08, 3.80, 2.62, "A reusable methodology",
     ["A generalisation-focused evaluation design that responds directly to documented "
      "benchmark-overfitting concerns.",
      "",
      "Applicable beyond this test case, to any domain-adaptation study."],
     accent=ACC2, body_size=12.5)
callout(s, M, TOP + 2.94, CW, 1.14, "What we deliberately do not claim",
        "We do not claim the negative-strategy comparison is new in general — it is well established for "
        "natural-language retrieval. We claim it is unrun for code, and that its interaction with "
        "generalisation has not been measured in any domain.",
        accent=WARN)

# ---------------------------------------------------------------- DIVIDER 5
divider("05", "Methodology",
        "Design, data, analysis and tools — with one variable and everything else held constant.")

# ---------------------------------------------------------------- 14  DESIGN
s = content_slide("Research Design", "05  ·  Methodology", "Methodology",
                  "Quantitative, controlled, and deliberately narrow.")
bullet(s, M, TOP + 0.04, CW, "Design:",
       "Quantitative, controlled experimental design. Four conditions compared under identical settings.",
       size=14)
bullet(s, M, TOP + 0.58, CW, "Control:",
       "Only the negative-sampling method differs between the three fine-tuned conditions. Model, loss, "
       "positive pairs, subset size, epochs, batch size and learning rate are held constant.", size=14)

rows = [
    ["Condition", "Fine-tuning", "Negative examples are chosen by", "Role in the comparison"],
    ["C0", "None — off the shelf", "n/a", "The control. Is fine-tuning worth it at all?"],
    ["C1", "Yes", "Random selection among non-clones", "The default practitioners actually use"],
    ["C2", "Yes", "BM25 retrieval — keyword-similar", "Hard negatives, cheap and deterministic"],
    ["C3", "Yes", "Base-model embeddings — semantically similar", "Hardest negatives, and the riskiest"],
]
table(s, M, TOP + 1.30, CW, 5, 4, rows, col_w=[0.62, 1.05, 1.85, 1.75], row_h=0.52, head_h=0.46,
      fsize=12, hsize=11.5, align=["c", "l", "l", "l"])
rect(s, M, TOP + 3.88, CW, 1.34, fill=ACC_SOFT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.06)
rect(s, M, TOP + 3.88, 0.06, 1.34, fill=ACC)
tf = tb(s, M + 0.3, TOP + 4.06, CW - 0.6, 1.02)
run(para(tf, first=True, space_after=4, line=1.2),
    "One variable, held constant everywhere else", size=13.5, bold=True, color=INK)
run(para(tf, space_after=0, line=1.2),
    "The same positive pairs are used in all four conditions. Because only the negatives change, any "
    "difference in the result is attributable to the negative-selection method.",
    size=12.5, color=BODY)

# ---------------------------------------------------------------- 15  NEGATIVES
s = content_slide("Building the Negative Pairs", "05  ·  Methodology", "Methodology",
                  "Three strategies, one anchor — worked on a concrete example.")
card(s, M, TOP + 0.06, 3.80, 1.62, "Random",
     ["Any non-clone snippet, chosen at random.", "Cost: none.  Difficulty: easy."],
     accent=MUTED, body_size=12)
card(s, M + 4.04, TOP + 0.06, 3.80, 1.62, "BM25  (keyword-similar)",
     ["BM25 index over the corpus; top keyword matches; known clones removed.",
      "Finds: shared vocabulary."],
     accent=ACC, body_size=12)
card(s, M + 8.09, TOP + 0.06, 3.80, 1.62, "Semantic  (model-similar)",
     ["Base model embeds the corpus once; top cosine matches; clones removed.",
      "Finds: what the model already confuses."],
     accent=ACC2, body_size=12)

code_panel(s, M, TOP + 1.86, 3.80, 1.92, "Anchor",
           ["void bubbleSort(int[] a) {",
            "  for (int i = 0; i < a.length; i++)",
            "    for (int j = 0; j < a.length-1; j++)",
            "      if (a[j] > a[j+1]) swap(a, j, j+1);",
            "}"])
code_panel(s, M + 4.04, TOP + 1.86, 3.80, 1.92, "Random negative  —  easy",
           ["// unrelated: no shared keywords",
            "Connection connectDb(String url) {",
            "  return DriverManager",
            "      .getConnection(url);",
            "}"])
code_panel(s, M + 8.09, TOP + 1.86, 3.80, 1.92, "BM25 negative  —  harder",
           ["// shares int[], a, for, swap",
            "void selectionSort(int[] a) {",
            "  for (int i = 0; i < a.length; i++)",
            "    swap(a, i, minIndex(a, i));",
            "}"])
callout(s, M, TOP + 3.90, CW, 1.06,
        "Verification step — build and test this first",
        ["Before a mined snippet is used as a negative, exclude every snippet that is a known clone "
         "of the anchor. Necessary for correctness — and, as the next slide shows, not sufficient.",
         "Look again at the BM25 example: a different sort. Is that really a non-clone?"],
        accent=WARN)

# ---------------------------------------------------------------- 16  FALSE NEGATIVES
s = content_slide("The Trap: Not Every “Non-clone” Is Really a Non-clone",
                  "05  ·  Methodology", "Methodology")
bullet(s, M, TOP + 0.04, CW, "The mechanism:",
       "The more similar a snippet you mine, the more likely it is a genuine clone that the dataset "
       "forgot to label. Harder negatives carry a larger false-negative burden.", size=13.5)
bullet(s, M, TOP + 0.62, CW, "Why the labels are incomplete:",
       "BigCloneBench records clone pairs per functionality and makes no claim about pairs it never "
       "labelled — within or across functionalities. Its own audit work estimates at least 15% of "
       "labels as subjective or erroneous.", size=13.5)
bullet(s, M, TOP + 1.26, CW, "What that would do:",
       "We would be training the model to push apart code that is genuinely equivalent — and the "
       "hardest condition would suffer most, for the wrong reason.", size=13.5)

# hardness vs contamination schematic
by = TOP + 1.98
rect(s, M, by, CW, 1.66, fill=RGBColor(0xF7, 0xFA, 0xFD), line=RULE, lw=0.9,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.05)
tf = tb(s, M + 0.3, by + 0.14, 5.0, 0.3)
run(para(tf, first=True), "HARDNESS AND CONTAMINATION RISE TOGETHER", size=10.5, bold=True, color=ACC)
bars = [("Random", 0.22, 0.10, MUTED), ("BM25", 0.55, 0.34, ACC), ("Semantic", 0.86, 0.62, ACC2)]
for i, (nm, hard, cont, col) in enumerate(bars):
    bx = M + 0.35 + i * 3.95
    tf = tb(s, bx, by + 0.56, 3.6, 0.28)
    run(para(tf, first=True), nm, size=12.5, bold=True, color=INK)
    rect(s, bx, by + 0.90, 3.6, 0.20, fill=RGBColor(0xE4, 0xEA, 0xF0))
    rect(s, bx, by + 0.90, 3.6 * hard, 0.20, fill=col)
    tf = tb(s, bx, by + 1.16, 3.6, 0.24)
    run(para(tf, first=True), "hardness", size=9.5, color=MUTED)
    rect(s, bx, by + 1.44, 3.6, 0.14, fill=RGBColor(0xF6, 0xE3, 0xE0))
    rect(s, bx, by + 1.44, 3.6 * cont, 0.14, fill=RGBColor(0xD9, 0x6A, 0x4A))
tf = tb(s, M + 0.3, by + 1.44 - 0.30, 0.05, 0.2)
run(para(tf, first=True), "", size=8)

callout(s, M, TOP + 3.82, CW, 1.24, "So we measure it — this is what makes the result interpretable",
        ["Sample 50 mined negatives from the BM25 condition and 50 from the semantic condition, and "
         "judge by eye whether each pair implements the same functionality.",
         "Without this, a “harder negatives are worse” result has two competing explanations: hardness "
         "genuinely backfires, or the hard set is simply more contaminated. The measurement separates them."],
        accent=ACC2)

# ---------------------------------------------------------------- 17  DATA COLLECTION
s = content_slide("Data Collection Methods", "05  ·  Methodology", "Methodology")
card(s, M, TOP + 0.06, 5.80, 1.98, "Source dataset",
     ["BigCloneBench, CodeXGLUE-processed version — real, labelled Java code-clone pairs.",
      "",
      "BigCloneBench proper: 8,915,130 true clone pairs and 288,367 false clone pairs across 43 "
      "functionalities, drawn from roughly 25,000 open-source projects.",
      "",
      "What we actually load: 9,134 Java fragments — 901,028 train / 415,416 validation / 415,416 test "
      "pairs."],
     accent=ACC, body_size=12.5)
card(s, M + 6.09, TOP + 0.06, 5.80, 1.98, "Positive pairs",
     ["True clone pairs from the dataset, used unchanged across all four conditions.",
      "",
      "Positives are identical between conditions by construction — the negative is the only thing "
      "that moves.",
      "",
      "A fixed subset is sampled once and reused, so no condition sees data the others do not."],
     accent=GOOD, body_size=12.5)

rows = [
    ["Negative strategy", "How the snippet is found", "Filter applied before use"],
    ["Random", "Any non-clone snippet, sampled at random", "Exclude known clones of the anchor"],
    ["BM25", "rank_bm25 index over the corpus; top keyword matches", "Exclude known clones of the anchor"],
    ["Semantic", "Base model embeds the corpus once; top cosine matches", "Exclude known clones of the anchor"],
]
table(s, M, TOP + 2.22, CW, 4, 3, rows, col_w=[0.85, 2.05, 1.35], row_h=0.48, head_h=0.44,
      fsize=12, hsize=11.5)
callout(s, M, TOP + 4.36, CW, 0.88, "On subset size",
        "Training uses a manageable fixed subset rather than all 901,028 pairs, for tractability. This "
        "is standard practice for a project of this length, and CodeXGLUE's own reference pipeline uses "
        "10% of the training data.",
        accent=ACC)

# ---------------------------------------------------------------- 18  FLOWCHART
s = content_slide("Research Process Flowchart", "05  ·  Methodology", "Methodology",
                  "The whole method on one slide.")
FX, FW = M, CW
rows_f = [
    ("single", "Domain dataset — BigCloneBench (CodeXGLUE, 9,134 Java fragments)", NAVY, WHITE, 0.52),
    ("single", "Positive pairs — fixed, shared by every condition", RGBColor(0x1F, 0x6F, 0xB2), WHITE, 0.52),
    ("triple", ["Random negatives", "BM25 negatives", "Semantic negatives"],
     [MUTED, ACC, ACC2], WHITE, 0.62),
    ("triple", ["Fine-tune  C1", "Fine-tune  C2", "Fine-tune  C3"],
     [MUTED, ACC, ACC2], WHITE, 0.58),
    ("single", "+  Baseline  C0  (no fine-tuning)  —  the control", RGBColor(0x4A, 0x55, 0x60), WHITE, 0.50),
    ("triple", ["Standard benchmark evaluation", "Generalisation evaluation (unseen functionalities)",
                "Embedding visualisation"],
     [ACC, ACC2, GOOD], WHITE, 0.66),
    ("single", "Results comparison across the four conditions", NAVY, WHITE, 0.50),
]
y = TOP + 0.06
prev_y = None
for kind, spec, fillc, txtc, hh in rows_f:
    if kind == "single":
        rect(s, FX + CW * 0.16, y, CW * 0.68, hh, fill=fillc,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.10)
        tf = tb(s, FX + CW * 0.16, y + 0.06, CW * 0.68, hh - 0.08, anchor=MSO_ANCHOR.MIDDLE)
        run(para(tf, first=True, align=PP_ALIGN.CENTER, space_after=0), spec, size=12.5,
            bold=True, color=txtc)
    else:
        bw = CW * 0.30
        gap = (CW - 3 * bw) / 2
        for i in range(3):
            bx = FX + i * (bw + gap)
            rect(s, bx, y, bw, hh, fill=fillc[i], shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.09)
            tf = tb(s, bx, y + 0.05, bw, hh - 0.06, anchor=MSO_ANCHOR.MIDDLE)
            run(para(tf, first=True, align=PP_ALIGN.CENTER, space_after=0, line=1.12), spec[i],
                size=12, bold=True, color=txtc)
    if prev_y is not None:
        arrow_down(s, FX + CW / 2, prev_y, y - prev_y, color=RGBColor(0xB8, 0xC2, 0xCC), w=0.15)
    prev_y = y + hh
    y = y + hh + 0.30

# ---------------------------------------------------------------- 19  DATA ANALYSIS
s = content_slide("Data Analysis Techniques", "05  ·  Methodology", "Methodology")
card(s, M, TOP + 0.06, 5.80, 2.06, "Quantitative comparison",
     ["Primary: F1, with precision and recall reported alongside — the metric CodeXGLUE reports, so "
      "our numbers are comparable to published work.",
      "",
      "Secondary: MAP@R — threshold-free, so it shows whether the F1 ordering is an artefact of "
      "threshold choice.",
      "",
      "The cosine threshold is selected on the validation split and reported."],
     accent=ACC, body_size=12.5)
card(s, M + 6.09, TOP + 0.06, 5.80, 2.06, "Generalisation analysis",
     ["Standard test split versus held-out, unseen functionality categories.",
      "",
      "Reported as  F1 seen  −  F1 unseen  =  Δ, per condition.",
      "",
      "Δ is the dependent variable: it is the gap itself we are testing, not the raw score."],
     accent=ACC2, body_size=12.5)
card(s, M, TOP + 2.30, 5.80, 1.86, "Visual and qualitative analysis",
     ["t-SNE / UMAP projection to compare embedding cluster separation across the four conditions.",
      "",
      "Fixed random seed and shared sample across all panels, so the comparison is fair."],
     accent=GOOD, body_size=12.5)
card(s, M + 6.09, TOP + 2.30, 5.80, 1.86, "Analysis discipline",
     ["Report differences between conditions, not absolute levels — the benchmark's labels are "
      "themselves imperfect.",
      "",
      "If F1 and MAP@R disagree about the ordering, report the disagreement rather than tuning it away."],
     accent=WARN, body_size=12.5)
callout(s, M, TOP + 4.34, CW, 0.94, "A caution that is written into the method, not added afterwards",
        "A 2D projection is an illustration, not evidence. Cluster shape depends on projection "
        "hyperparameters, so the visualisation is presented after the results table and never in place of it.",
        accent=WARN)

# ---------------------------------------------------------------- 20  TOOLS
s = content_slide("Tools and Resources Used", "05  ·  Methodology", "Methodology")
tools = [
    ("Languages", ["Python", "PyTorch"], ACC),
    ("Modelling", ["Hugging Face Transformers", "sentence-transformers"], ACC),
    ("Base model", ["GraphCodeBERT", "(pretrained, 768-d)"], GOOD),
    ("Negative mining", ["rank_bm25 — BM25 index", "Corpus encoded once for semantic mining"], ACC2),
    ("Dataset", ["BigCloneBench / CodeXGLUE", "901k / 415k / 415k pairs"], ACC),
    ("Visualisation", ["t-SNE — scikit-learn", "UMAP — umap-learn"], GOOD),
]
cw2 = (CW - 2 * 0.26) / 3
for i, (cat, items, col) in enumerate(tools):
    r, c = divmod(i, 3)
    x = M + c * (cw2 + 0.26)
    y = TOP + 0.08 + r * (1.72 + 0.26)
    rect(s, x, y, cw2, 1.72, fill=PANEL, line=RULE, lw=0.9,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.06)
    rect(s, x, y, 0.055, 1.72, fill=col)
    tf = tb(s, x + 0.24, y + 0.16, cw2 - 0.44, 0.32)
    run(para(tf, first=True, space_after=4), cat.upper(), size=10.5, bold=True, color=col)
    tf = tb(s, x + 0.24, y + 0.52, cw2 - 0.44, 1.0)
    for j, it in enumerate(items):
        run(para(tf, first=(j == 0), space_after=5, line=1.15), it, size=12.5, color=BODY)

rect(s, M, TOP + 3.94, CW, 1.10, fill=ACC_SOFT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.06)
rect(s, M, TOP + 3.94, 0.06, 1.10, fill=ACC)
tf = tb(s, M + 0.3, TOP + 4.10, CW - 0.6, 0.82)
run(para(tf, first=True, space_after=4, line=1.2), "Compute", size=13, bold=True, color=INK)
run(para(tf, space_after=0, line=1.2),
    "Free-tier GPU (Google Colab). Feasible because the corpus is 9,134 fragments, so negative mining "
    "and retrieval are single matrix operations — no approximate-index infrastructure is required.",
    size=12.5, color=BODY)

# ---------------------------------------------------------------- DIVIDER 6
divider("06", "Proposed Outcomes",
        "What we expect to find, why it matters, and where it leads next.")

# ---------------------------------------------------------------- 21  EXPECTED RESULTS
s = content_slide("Expected Results", "06  ·  Proposed Outcomes", "Proposed Outcomes")
fig_path = os.path.join(ASSETS, "hypotheses.png")
make_hypotheses_figure(fig_path)
s.shapes.add_picture(fig_path, E(M), E(TOP + 0.06), height=E(2.72))

RX = M + 7.55
RW = CW - 7.55
bullet(s, RX, TOP + 0.06, RW, "Fine-tuning:",
       "At least one strategy is expected to beat the off-the-shelf baseline.", size=12.5, space_after=8)
bullet(s, RX, TOP + 0.76, RW, "Random vs the rest:",
       "Random is expected to trail both mined strategies.", size=12.5, space_after=8)
bullet(s, RX, TOP + 1.46, RW, "BM25 vs semantic:",
       "Genuinely open — published comparisons put them within noise.", size=12.5, space_after=8)
bullet(s, RX, TOP + 2.16, RW, "Generalisation:",
       "A drop is expected. The new finding is whether the drop depends on the strategy.",
       size=12.5, space_after=8)

callout(s, M, TOP + 2.94, CW, 1.24,
        "Which hypothesis the literature actually supports",
        "The literature-supported prediction is the inverted-U, not a straight line: harder negatives "
        "help up to a point, and mining too aggressively introduces false negatives that undo the gain. "
        "All three shapes are legitimate outcomes, and the deck is built to present whichever occurs.",
        accent=ACC2)

# ---------------------------------------------------------------- 22  IMPACT
s = content_slide("Potential Impact on the Field", "06  ·  Proposed Outcomes", "Proposed Outcomes")
card(s, M, TOP + 0.10, 5.80, 2.34, "Guidance that transfers",
     ["Evidence-based guidance for selecting a negative-sampling strategy when adapting embeddings to "
      "any specialised domain — not limited to code.",
      "",
      "The question is a method question. A method answer generalises to legal text, support tickets, "
      "clinical notes and bug reports, even though the experiment runs on Java."],
     accent=ACC, body_size=12.5)
card(s, M + 6.09, TOP + 0.10, 5.80, 2.34, "A stronger evaluation habit",
     ["Reinforces generalisation testing as standard practice in domain-adaptation research, rather "
      "than relying on a single standard benchmark split.",
      "",
      "A held-out category split costs one extra experiment and answers a question the standard split "
      "cannot."],
     accent=GOOD, body_size=12.5)
bullet(s, M, TOP + 2.66, CW, "For practitioners:",
       "A costed answer to “should we build the mining pipeline, or is random sampling fine?” — "
       "currently answered by intuition.", size=13.5)
bullet(s, M, TOP + 3.24, CW, "For researchers:",
       "A clean, reproducible four-condition comparison with a published benchmark to check against, "
       "and a measured false-negative rate rather than an assumed one.", size=13.5)
callout(s, M, TOP + 3.92, CW, 0.92, "Framing the impact honestly",
        "The likely size of any effect is modest. The value of the work is that it converts a choice "
        "everyone makes informally into a choice with evidence behind it.",
        accent=ACC)

# ---------------------------------------------------------------- 23  FUTURE DIRECTIONS
s = content_slide("Future Research Directions", "06  ·  Proposed Outcomes", "Proposed Outcomes")
card(s, M, TOP + 0.10, 3.80, 2.30, "Other domains",
     ["Apply the same negative-sampling comparison to legal text, duplicate bug reports and medical "
      "record matching.",
      "",
      "Tests whether any finding is a property of code or a property of the method."],
     accent=ACC, body_size=12.5)
card(s, M + 4.04, TOP + 0.10, 3.80, 2.30, "Automatic strategy selection",
     ["Choose the negative-sampling strategy for a new, unseen domain automatically, using cheap pilot "
      "runs instead of the user's judgement.",
      "",
      "Deliberately excluded here, to keep the comparison clean and interpretable."],
     accent=GOOD, body_size=12.5)
card(s, M + 8.09, TOP + 0.10, 3.80, 2.30, "Hybrid mining",
     ["Combine random, BM25 and semantic negatives in weighted proportions, and sweep the mixing ratio.",
      "",
      "The natural follow-up once the three pure conditions are measured."],
     accent=ACC2, body_size=12.5)
callout(s, M, TOP + 2.62, CW, 1.10, "One thing worth knowing before starting the hybrid direction",
        "A fused BM25 + semantic strategy has already been tried in natural-language retrieval and lost "
        "on four of five tasks. It remains a reasonable next experiment here — but the prior is weaker "
        "than it looks.",
        accent=WARN)
callout(s, M, TOP + 3.92, CW, 0.92, "Also worth doing",
        "Refresh the mined negatives part-way through training, so the hard examples track the model as "
        "it learns rather than staying fixed from epoch zero.",
        accent=ACC)

# ---------------------------------------------------------------- DIVIDER 7
divider("07", "Timeline of the Project",
        "Initiation → planning → execution → closure, with the research tasks mapped inside.")

# ---------------------------------------------------------------- 24  TIMELINE
s = content_slide("Timeline of the Project", "07  ·  Timeline", "Timeline",
                  "EL phase structure, with the four-week research execution mapped into it.")
phases = [
    ("Initiation", "Weeks 1–2", ACC,
     ["Define project goals and objectives",
      "Identify resources and compute",
      "Feasibility: dataset access, tooling",
      "Confirm scope and stop condition"]),
    ("Planning", "Weeks 3–4", ACC,
     ["Detailed project plan and task list",
      "Pipeline design and data contracts",
      "Negative-mining verification built and tested",
      "Fix subset size from a throughput test"]),
    ("Execution", "Weeks 5–10", ACC2,
     ["Baseline C0 evaluated first",
      "C1 / C2 / C3 fine-tuning runs",
      "Standard benchmark evaluation",
      "Generalisation test · visualisation · demo"]),
    ("Closure", "Weeks 11–13", GOOD,
     ["Final results assembled into one table",
      "Write-up, limitations and EL deliverable",
      "Project review presentation",
      "Lessons learned documented"]),
]
pw = (CW - 3 * 0.24) / 4
for i, (name, wk, col, items) in enumerate(phases):
    x = M + i * (pw + 0.24)
    rect(s, x, TOP + 0.06, pw, 0.62, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.10)
    tf = tb(s, x, TOP + 0.13, pw, 0.5)
    run(para(tf, first=True, align=PP_ALIGN.CENTER, space_after=1), name, size=14, bold=True, color=WHITE)
    run(para(tf, align=PP_ALIGN.CENTER, space_after=0), wk, size=11, color=RGBColor(0xE3, 0xEC, 0xF5))
    rect(s, x, TOP + 0.82, pw, 2.52, fill=PANEL, line=RULE, lw=0.9,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.05)
    tf = tb(s, x + 0.20, TOP + 0.98, pw - 0.40, 2.24)
    for j, it in enumerate(items):
        p = para(tf, first=(j == 0), space_after=8, line=1.15)
        run(p, "·  ", size=11.5, bold=True, color=col)
        run(p, it, size=11.5, color=BODY)

rect(s, M, TOP + 3.46, CW, 1.02, fill=ACC_SOFT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.05)
rect(s, M, TOP + 3.46, 0.06, 1.02, fill=ACC2)
tf = tb(s, M + 0.3, TOP + 3.62, CW - 0.6, 0.72)
run(para(tf, first=True, space_after=5, line=1.2),
    "The research execution, in four working weeks", size=13, bold=True, color=INK)
run(para(tf, space_after=0, line=1.2),
    "Week 1  data pipeline, negative mining verified, one end-to-end run   ·   "
    "Week 2  the three conditions trained and evaluated on the standard split   ·   "
    "Week 3  generalisation test, visualisation, demo   ·   "
    "Week 4  results, write-up, limitations, presentation",
    size=12, color=BODY)
callout(s, M, TOP + 4.64, CW, 0.88, "Built-in fallback",
        "If time compresses, drop to two strategies — random and semantic — which still answers whether "
        "the strategy matters. The cost is that a peaked (inverted-U) result can no longer be "
        "distinguished from a flat one.",
        accent=WARN)

# ---------------------------------------------------------------- DIVIDER 8
divider("08", "Conclusion",
        "What the project is, in three points, and what happens next.")

# ---------------------------------------------------------------- 25  CONCLUSION
s = content_slide("Conclusion", "08  ·  Conclusion", "Conclusion")
bullet(s, M, TOP + 0.06, CW, "The gap:",
       "Generic embeddings may not capture domain-specific similarity. Contrastive fine-tuning is a "
       "known fix — but how the negative examples are chosen remains under-studied.", size=13.5)
bullet(s, M, TOP + 0.68, CW, "The test case:",
       "Code similarity detection, using GraphCodeBERT and BigCloneBench — a pairing chosen so our "
       "numbers sit beside published ones.", size=13.5)
bullet(s, M, TOP + 1.30, CW, "The design:",
       "Negative-sampling strategy as a controlled variable, with generalisation testing built in from "
       "the start rather than added as a critique.", size=13.5)

rect(s, M, TOP + 2.06, CW, 1.16, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.05)
rect(s, M, TOP + 2.06, 0.07, 1.16, fill=ACC2)
tf = tb(s, M + 0.42, TOP + 2.26, CW - 0.85, 0.85)
run(para(tf, first=True, space_after=5), "Final thought", size=10.5, bold=True, color=ACC2)
run(para(tf, space_after=0, line=1.24),
    "“How we train a model to recognize similarity is as important as the architecture itself — this "
    "project treats the training methodology as a genuine subject of investigation, not an "
    "implementation detail.”",
    size=14, color=WHITE)

tf = tb(s, M, TOP + 3.44, CW, 0.32)
run(para(tf, first=True), "Next steps", size=14.5, bold=True, color=INK)
steps = [
    ("1", "Implement and validate the three negative-mining pipelines, with the clone-exclusion check tested"),
    ("2", "Run comparative training and generalisation evaluation across all four conditions"),
    ("3", "Extend the comparison to a second domain, to test whether the finding transfers beyond code"),
]
for i, (n, txt) in enumerate(steps):
    y = TOP + 3.88 + i * 0.56
    rect(s, M, y, CW, 0.48, fill=PANEL, line=RULE, lw=0.9,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.10)
    chip(s, M + 0.16, y + 0.07, 0.34, 0.34, n, fill=ACC, size=11.5)
    tf = tb(s, M + 0.66, y + 0.07, CW - 0.86, 0.38)
    run(para(tf, first=True, space_after=0, line=1.15), txt, size=12.5, color=BODY)

# ---------------------------------------------------------------- 26  REFERENCES
s = content_slide("References", "09  ·  References", "References")
refs_l = [
    "[1]  D. Guo et al., “GraphCodeBERT: Pre-training Code Representations with Data Flow,” in Proc. Int. Conf. Learning Representations (ICLR), 2021.",
    "[2]  P. Jain, A. Jain, T. Zhang, P. Abbeel, J. Gonzalez, and I. Stoica, “Contrastive Code Representation Learning,” in Proc. Conf. Empirical Methods in Natural Language Processing (EMNLP), 2021, pp. 5954–5971.",
    "[3]  T. Sonnekalb, B. Gruner, C.-A. Brust, and P. Mader, “Generalizability of Code Clone Detection on CodeBERT,” in Proc. 37th IEEE/ACM Int. Conf. Automated Software Engineering (ASE), Rochester, MI, USA, 2022.",
    "[4]  S. Liu, B. Wu, X. Xie, G. Meng, and Y. Liu, “ContraBERT: Enhancing Code Pre-trained Models via Contrastive Learning,” in Proc. IEEE/ACM 45th Int. Conf. Software Engineering (ICSE), 2023, pp. 2476–2487.",
    "[5]  M. Khajezade, F. H. Fard, and M. S. Shehata, “Evaluating Few-shot and Contrastive Learning Methods for Code Clone Detection,” Empirical Software Engineering, vol. 29, no. 6, p. 163, 2024.",
]
refs_r = [
    "[6]  P. H. Le-Khac, G. Healy, and A. F. Smeaton, “Contrastive Representation Learning: A Framework and Review,” IEEE Access, vol. 8, pp. 193907–193934, 2020.",
    "[7]  K. Kitsios, F. Sovrano, E. T. Barr, and A. Bacchelli, “Detecting Semantic Clones of Unseen Functionality,” in Proc. 40th IEEE/ACM Int. Conf. Automated Software Engineering (ASE), 2025.",
    "[8]  N. Thakur, N. Reimers, J. Daxenberger, and I. Gurevych, “Augmented SBERT: Data Augmentation Method for Improving Bi-Encoders for Pairwise Sentence Scoring Tasks,” in Proc. NAACL-HLT, 2021.",
    "[9]  J. Zhan, J. Mao, Y. Liu, J. Guo, M. Zhang, and S. Ma, “Optimizing Dense Retrieval Model Training with Hard Negatives,” in Proc. 44th Int. ACM SIGIR Conf. Research and Development in Information Retrieval (SIGIR), 2021.",
    "[10]  J. Krinke, “BigCloneBench Considered Harmful for Machine Learning,” in Proc. 10th Int. Workshop on Software Clones (IWSC), 2022.",
]
for col_i, refs in enumerate([refs_l, refs_r]):
    x = M + col_i * (CW / 2 + 0.18)
    w = CW / 2 - 0.18
    tf = tb(s, x, TOP + 0.06, w, 5.3)
    for i, r in enumerate(refs):
        run(para(tf, first=(i == 0), space_after=11, line=1.16), r, size=10.5, color=BODY)

# ---------------------------------------------------------------- 27  EL DELIVERABLE
s = content_slide("Expected Outcome by the End of First Year", "EL  ·  Deliverable", "EL Deliverable",
                  "EL work is expected to produce one of the following tangible outputs.")
opts = [
    ("Patent Filing", "Not the primary target", ACC2,
     ["The contribution is a comparative empirical study, not a novel functional prototype or process, "
      "so patentability is limited. Not pursued as the main outcome."]),
    ("Journal Publication", "Strong fit", GOOD,
     ["A paper detailing the four-condition methodology, experimental design and results is well suited "
      "to a software-engineering or applied-ML journal, following the format of several reviewed "
      "studies (e.g. Empirical Software Engineering)."]),
    ("Conference Publication", "Primary planned outcome", ACC,
     ["A full paper in IEEE format covering the comparative methodology, the findings across the three "
      "negative-sampling strategies, and the generalisation evaluation, submitted to a relevant "
      "software-engineering or AI conference."]),
    ("Research Proposal", "Natural extension", ACC,
     ["Propose applying the same comparison to further domains beyond code, and explore automatic "
      "negative-strategy selection for new domains as a larger follow-up project."]),
]
ow = (CW - 3 * 0.22) / 4
for i, (name, tag, col, body) in enumerate(opts):
    x = M + i * (ow + 0.22)
    rect(s, x, TOP + 0.08, ow, 3.06, fill=PANEL, line=RULE, lw=0.9,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.05)
    rect(s, x, TOP + 0.08, ow, 0.075, fill=col)
    tf = tb(s, x + 0.18, TOP + 0.28, ow - 0.36, 0.5)
    run(para(tf, first=True, space_after=3), name, size=13, bold=True, color=INK)
    chip(s, x + 0.18, TOP + 0.80, ow - 0.36, 0.32, tag, fill=col, size=10)
    tf = tb(s, x + 0.18, TOP + 1.26, ow - 0.36, 1.7)
    for j, ln in enumerate(body):
        run(para(tf, first=(j == 0), space_after=6, line=1.16), ln, size=11, color=BODY)

rect(s, M, TOP + 3.36, CW, 1.00, fill=ACC_SOFT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.06)
rect(s, M, TOP + 3.36, 0.06, 1.00, fill=ACC)
tf = tb(s, M + 0.3, TOP + 3.52, CW - 0.6, 0.74)
run(para(tf, first=True, space_after=4, line=1.2),
    "Planned focus for this EL", size=13, bold=True, color=INK)
run(para(tf, space_after=0, line=1.2),
    "Conference publication in IEEE format, with the underlying study structured to also support a "
    "future journal submission or a research-proposal extension.",
    size=12, color=BODY)
plain(s, M, 5.90, CW, "Date on slide: 17 September 2026", size=10.5, color=MUTED)

# ---------------------------------------------------------------- 28  CLOSING (image only)
PAGE["n"] += 1
s = prs.slides.add_slide(BLANK)
closing = os.path.join(ASSETS, "closing-embedding-space.png")
if os.path.exists(closing):
    s.shapes.add_picture(closing, E(0), E(0), width=E(W), height=E(H))
else:
    rect(s, 0, 0, W, H, fill=NAVY)

prs.save(OUT)
print(f"saved {OUT}")
print(f"slides: {len(prs.slides.__iter__.__self__._sldIdLst)}")
